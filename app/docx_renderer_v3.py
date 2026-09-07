from __future__ import annotations

import io
import math
import re
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Pt

EMU_PER_PT = 12700
WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
PICTURE_URI = "http://schemas.openxmlformats.org/drawingml/2006/picture"
nsmap.setdefault("wps", WPS_URI)


def el(tag: str):
    return OxmlElement(tag)


def _anchor_position(x: float, y: float, w: float, h: float, ident: int, name: str, behind: bool = False):
    anchor = el("wp:anchor")
    for k, v in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "1" if behind else "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(k, v)
    simple = el("wp:simplePos")
    simple.set("x", "0"); simple.set("y", "0")
    anchor.append(simple)
    pos_h = el("wp:positionH"); pos_h.set("relativeFrom", "page")
    off_h = el("wp:posOffset"); off_h.text = str(int(x * EMU_PER_PT)); pos_h.append(off_h); anchor.append(pos_h)
    pos_v = el("wp:positionV"); pos_v.set("relativeFrom", "page")
    off_v = el("wp:posOffset"); off_v.text = str(int(y * EMU_PER_PT)); pos_v.append(off_v); anchor.append(pos_v)
    extent = el("wp:extent")
    extent.set("cx", str(max(1, int(w * EMU_PER_PT))))
    extent.set("cy", str(max(1, int(h * EMU_PER_PT))))
    anchor.append(extent)
    effect = el("wp:effectExtent")
    for k in ("l", "t", "r", "b"): effect.set(k, "0")
    anchor.append(effect)
    anchor.append(el("wp:wrapNone"))
    docpr = el("wp:docPr")
    docpr.set("id", str(ident)); docpr.set("name", name)
    anchor.append(docpr)
    cNv = el("wp:cNvGraphicFramePr")
    cNv.append(el("a:graphicFrameLocks"))
    anchor.append(cNv)
    return anchor


def make_textbox(x: float, y: float, w: float, h: float, text: str, font_pt: float,
                 bold: bool = False, italic: bool = False, align: str = "left", ident: int = 1):
    anchor = _anchor_position(x, y, w, h, ident, f"Text {ident}")
    graphic = el("a:graphic")
    data = el("a:graphicData"); data.set("uri", WPS_URI)
    wsp = el("wps:wsp")
    wsp.append(el("wps:cNvSpPr"))
    sppr = el("wps:spPr")
    xfrm = el("a:xfrm")
    off = el("a:off"); off.set("x", "0"); off.set("y", "0")
    ext = el("a:ext"); ext.set("cx", str(max(1, int(w * EMU_PER_PT)))); ext.set("cy", str(max(1, int(h * EMU_PER_PT))))
    xfrm.extend([off, ext]); sppr.append(xfrm)
    geom = el("a:prstGeom"); geom.set("prst", "rect"); geom.append(el("a:avLst")); sppr.append(geom)
    sppr.append(el("a:noFill")); wsp.append(sppr)

    txbx = el("wps:txbx")
    content = el("w:txbxContent")
    lines = str(text).splitlines() or [""]
    p = el("w:p")
    ppr = el("w:pPr")
    spacing = el("w:spacing")
    spacing.set(qn("w:before"), "0"); spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240"); spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    jc = el("w:jc"); jc.set(qn("w:val"), align); ppr.append(jc)
    p.append(ppr)
    for i, line in enumerate(lines):
        if i:
            p.append(el("w:br"))
        r = el("w:r")
        rpr = el("w:rPr")
        sz = el("w:sz"); sz.set(qn("w:val"), str(max(8, min(96, round(font_pt * 2))))); rpr.append(sz)
        szcs = el("w:szCs"); szcs.set(qn("w:val"), str(max(8, min(96, round(font_pt * 2))))); rpr.append(szcs)
        if bold: rpr.append(el("w:b"))
        if italic: rpr.append(el("w:i"))
        r.append(rpr)
        wt = el("w:t"); wt.text = line; r.append(wt); p.append(r)
    content.append(p); txbx.append(content); wsp.append(txbx)
    body = el("wps:bodyPr")
    for k in ("lIns", "tIns", "rIns", "bIns"): body.set(k, "0")
    body.set("anchor", "t"); body.set("vert", "horz")
    wsp.append(body)
    data.append(wsp); graphic.append(data); anchor.append(graphic)
    return anchor


class V3Renderer:
    """Hybrid PDF->DOCX renderer with coordinate-preserved images and editable text."""

    def __init__(self, pdf_path: Path):
        self.pdf = fitz.open(pdf_path)
        self.ident = 2000

    def render(self, pages: list[dict], out_path: Path) -> None:
        doc = Document()
        if doc.paragraphs:
            p = doc.paragraphs[0]._element
            p.getparent().remove(p)
        for i, item in enumerate(pages):
            page = self.pdf[item["page_index"]]
            section = doc.sections[0] if i == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
            section.page_width = Pt(page.rect.width)
            section.page_height = Pt(page.rect.height)
            section.top_margin = section.bottom_margin = Pt(0)
            section.left_margin = section.right_margin = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
            self._render_page(doc, page, item)
        self.pdf.close()
        doc.save(out_path)

    def _render_page(self, doc, page, item):
        regions = item.get("regions", [])
        self._render_images(doc, page)
        for region in regions:
            label = str(region.get("label", "text")).lower()
            bbox = region.get("bbox") or [0, 0, page.rect.width, page.rect.height]
            lines = region.get("lines") or []
            if not lines:
                continue
            lines = sorted(lines, key=lambda x: (x["bbox"][1], x["bbox"][0]))
            if label == "table" and self._render_table(doc, lines, bbox):
                continue
            for block in self._group_lines(lines, label):
                self._render_block(doc, block, label, page.rect)

    def _render_images(self, doc, page):
        seen = set()
        for image in page.get_images(full=True):
            xref = image[0]
            rects = page.get_image_rects(xref)
            for rect in rects:
                key = (xref, round(rect.x0, 1), round(rect.y0, 1), round(rect.x1, 1), round(rect.y1, 1))
                if key in seen or rect.width <= 1 or rect.height <= 1:
                    continue
                seen.add(key)
                try:
                    pix = fitz.Pixmap(self.pdf, xref)
                    if pix.alpha:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    blob = pix.tobytes("png")
                    # Keep large scanned-page images as page backgrounds so OCR text remains visible.
                    area_ratio = (rect.width * rect.height) / max(1.0, page.rect.width * page.rect.height)
                    self._add_positioned_image(doc, blob, rect.x0, rect.y0, rect.width, rect.height,
                                               behind=area_ratio >= 0.80)
                except Exception:
                    continue

    def _add_positioned_image(self, doc, blob, x, y, w, h, behind=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        inline = run.add_picture(io.BytesIO(blob), width=Pt(max(1, w)), height=Pt(max(1, h)))
        old_inline = inline._inline
        pic = old_inline.find(qn("pic:pic"))
        if pic is None:
            return
        pic = pic.__deepcopy__({})
        old_inline.getparent().remove(old_inline)

        anchor = _anchor_position(x, y, w, h, self.ident, f"Image {self.ident}", behind=behind)
        self.ident += 1
        graphic = el("a:graphic")
        data = el("a:graphicData"); data.set("uri", PICTURE_URI)
        data.append(pic)
        graphic.append(data)
        anchor.append(graphic)
        p._p.append(el("w:drawing"))
        p._p[-1].append(anchor)

    def _group_lines(self, lines, label):
        groups = []
        for line in lines:
            b = line["bbox"]; h = max(1.0, b[3] - b[1])
            if not groups:
                groups.append([line]); continue
            prev = groups[-1]; pb = prev[-1]["bbox"]; ph = max(1.0, pb[3] - pb[1])
            gap = b[1] - pb[3]
            xdiff = abs(b[0] - pb[0])
            overlap = min(b[2], pb[2]) - max(b[0], pb[0])
            compatible_height = 0.55 <= h / ph <= 1.8
            same_column = xdiff <= max(8, ph * 1.8) or overlap > 0
            max_gap = max(4.0, ph * 0.9)
            if label in {"title", "section_header", "header", "footer", "caption", "figure_caption", "table_caption"}:
                compatible = gap <= ph * 0.35 and xdiff <= ph * 1.5
            else:
                compatible = gap <= max_gap and same_column and compatible_height
            if compatible: prev.append(line)
            else: groups.append([line])
        return groups

    def _render_block(self, doc, block, label, page_rect):
        x0 = min(float(x["bbox"][0]) for x in block); y0 = min(float(x["bbox"][1]) for x in block)
        x1 = max(float(x["bbox"][2]) for x in block); y1 = max(float(x["bbox"][3]) for x in block)
        text = "\n".join(str(x.get("text", "")).strip() for x in block if str(x.get("text", "")).strip())
        if not text: return
        font_pt = self._font_size(block)
        bold = label in {"title", "section_header", "header"} or font_pt >= 15
        italic = label in {"caption", "figure_caption", "table_caption"}
        align = self._alignment(block, page_rect.width)
        body_like = label in {"plain_text", "text", "paragraph", "list", "list_item", "body"}
        narrow = (x1 - x0) < page_rect.width * 0.88
        simple = all("\n" not in str(x.get("text", "")) for x in block)
        if body_like and narrow and simple and self._is_flow_safe(block):
            self._native_paragraph(doc, block, font_pt, bold, italic, align); return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        anchor = make_textbox(x0, y0, max(3, x1 - x0), max(font_pt * 1.25, y1 - y0 + 2), text,
                              font_pt, bold, italic, align, self.ident)
        self.ident += 1
        drawing = el("w:drawing"); drawing.append(anchor); p._p.append(drawing)

    def _native_paragraph(self, doc, block, font_pt, bold, italic, align):
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.05
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        for i, line in enumerate(block):
            if i: p.add_run().add_break()
            run = p.add_run(str(line.get("text", "")).strip())
            run.font.size = Pt(font_pt); run.bold = bold; run.italic = italic
            run.font.name = self._font_name(line.get("text", ""))
        x0 = min(float(x["bbox"][0]) for x in block)
        if x0 > 3: pf.left_indent = Pt(x0)

    def _render_table(self, doc, lines, bbox):
        rows = self._cluster_rows(lines)
        if len(rows) < 2: return False
        col_edges = self._infer_columns(rows, bbox)
        if len(col_edges) < 2: return False
        table = doc.add_table(rows=len(rows), cols=len(col_edges) - 1); table.autofit = False
        for r, row in enumerate(rows):
            for line in row:
                b = line["bbox"]; cx = (b[0] + b[2]) / 2
                c = min(len(col_edges) - 2, max(0, self._edge_index(col_edges, cx)))
                cell = table.cell(r, c); text = str(line.get("text", "")).strip()
                cell.text = (cell.text + " " + text).strip() if cell.text else text
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size = Pt(self._font_size([line]))
        return True

    @staticmethod
    def _cluster_rows(lines):
        rows = []
        for line in sorted(lines, key=lambda z: (z["bbox"][1], z["bbox"][0])):
            cy = (line["bbox"][1] + line["bbox"][3]) / 2; placed = False
            for row in rows:
                rcy = sum((x["bbox"][1] + x["bbox"][3]) / 2 for x in row) / len(row)
                rh = max(x["bbox"][3] - x["bbox"][1] for x in row)
                if abs(cy - rcy) <= max(3, rh * 0.65): row.append(line); placed = True; break
            if not placed: rows.append([line])
        return [sorted(r, key=lambda z: z["bbox"][0]) for r in rows]

    @staticmethod
    def _infer_columns(rows, bbox):
        xs = sorted(v for row in rows for line in row for v in (line["bbox"][0], line["bbox"][2]))
        clusters = []
        for x in xs:
            if not clusters or abs(x - clusters[-1][-1]) > 12: clusters.append([x])
            else: clusters[-1].append(x)
        edges = [sum(c) / len(c) for c in clusters]
        out = []
        for x in edges:
            if not out or x - out[-1] > 18: out.append(x)
        if len(out) < 2: return []
        return [max(float(bbox[0]), out[0])] + out[1:-1] + [min(float(bbox[2]), out[-1])]

    @staticmethod
    def _edge_index(edges, cx):
        for i in range(len(edges) - 1):
            if edges[i] <= cx <= edges[i + 1]: return i
        return len(edges) - 2

    @staticmethod
    def _is_flow_safe(block):
        return len(block) <= 12 and not any(len(str(x.get("text", ""))) > 180 for x in block)

    @staticmethod
    def _font_size(block):
        heights = sorted(max(1.0, float(x["bbox"][3]) - float(x["bbox"][1])) for x in block)
        if not heights: return 10.0
        h = heights[len(heights) // 2]
        return max(6.0, min(40.0, h * 0.78))

    @staticmethod
    def _alignment(block, page_width):
        x0 = min(float(x["bbox"][0]) for x in block); x1 = max(float(x["bbox"][2]) for x in block)
        center = (x0 + x1) / 2
        if abs(center - page_width / 2) < page_width * 0.08 and (x1 - x0) < page_width * 0.75: return "center"
        if x0 > page_width * 0.62: return "right"
        return "left"

    @staticmethod
    def _font_name(text):
        return "Noto Sans CJK SC" if re.search(r"[\u3400-\u9fff]", text) else "Aptos"


CoordinateDocxRenderer = V3Renderer
