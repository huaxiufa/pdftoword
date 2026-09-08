from __future__ import annotations

import math
import tempfile
from copy import deepcopy
from pathlib import Path

import fitz
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

EMU_PER_PT = 12700


def _remove_cell_margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")


def _remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _set_table_width(table, width):
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert(0, tblW)
    tblW.set(qn("w:w"), str(int(width.twips)))
    tblW.set(qn("w:type"), "dxa")
    table.columns[0].width = width
    table.cell(0, 0).width = width


def _set_run_font(run, span):
    font = span.get("font") or "Arial"
    size = float(span.get("size") or 10)
    run.font.name = font
    run.font.size = Pt(max(1, size))
    flags = int(span.get("flags") or 0)
    run.bold = bool(flags & 16)
    run.italic = bool(flags & 2)
    color = span.get("color")
    if isinstance(color, int):
        run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        rFonts.set(qn(f"w:{key}"), font)


def _image_bytes(doc, block):
    data = block.get("image")
    if data:
        try:
            if block.get("mask"):
                return fitz.Pixmap(fitz.Pixmap(data), fitz.Pixmap(block["mask"])).tobytes("png")
            return data
        except Exception:
            pass
    xref = int(block.get("xref") or 0)
    if xref:
        try:
            return doc.extract_image(xref)["image"]
        except Exception:
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.alpha or pix.n - pix.alpha > 3:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                return pix.tobytes("png")
            except Exception:
                pass
    return None


def _clip_image(path, rect, page_rect):
    clipped = rect & page_rect
    if clipped.is_empty or clipped == rect:
        return path, rect
    with Image.open(path) as image:
        image = image.convert("RGBA")
        sx = image.width / max(rect.width, 0.001)
        sy = image.height / max(rect.height, 0.001)
        box = (
            max(0, round((clipped.x0 - rect.x0) * sx)),
            max(0, round((clipped.y0 - rect.y0) * sy)),
            min(image.width, round((clipped.x1 - rect.x0) * sx)),
            min(image.height, round((clipped.y1 - rect.y0) * sy)),
        )
        target = path.with_name(path.stem + "-clip.png")
        image.crop(box).save(target, "PNG")
    return target, clipped


def _rotation(matrix):
    if not matrix:
        return 0
    try:
        angle = math.degrees(math.atan2(float(matrix[1]), float(matrix[0]))) % 360
    except Exception:
        return 0
    return min((0, 90, 180, 270), key=lambda x: abs(((angle - x + 180) % 360) - 180))


def _rotate_image(path, angle):
    if angle == 0:
        return path
    with Image.open(path) as image:
        target = path.with_name(path.stem + f"-rot{angle}.png")
        image.convert("RGBA").rotate(-angle, expand=True, resample=Image.Resampling.BICUBIC).save(target, "PNG")
    return target


def _add_anchor(cell, image_path, rect):
    width = max(1.0, rect.width)
    height = max(1.0, rect.height)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width / 72), height=Inches(height / 72))
    drawing = inline._inline

    anchor = OxmlElement("wp:anchor")
    for key in ("distT", "distB", "distL", "distR"):
        anchor.set(key, "0")
    anchor.set("behindDoc", "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    extent = OxmlElement("wp:extent")
    extent.set("cx", str(round(width * EMU_PER_PT)))
    extent.set("cy", str(round(height * EMU_PER_PT)))
    anchor.append(extent)

    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)

    for tag, coord in (("wp:positionH", rect.x0), ("wp:positionV", rect.y0)):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset")
        off.text = str(round(coord * EMU_PER_PT))
        pos.append(off)
        anchor.append(pos)

    anchor.append(OxmlElement("wp:wrapNone"))
    for child in list(drawing):
        anchor.append(deepcopy(child))
    drawing.getparent().replace(drawing, anchor)


def _extract_images(page, work_dir):
    work_dir.mkdir(parents=True, exist_ok=True)
    results = []
    seen = set()
    blocks = page.get_text("dict", sort=False).get("blocks", [])

    for index, block in enumerate(blocks):
        if block.get("type") != 1 or not block.get("bbox"):
            continue
        rect = fitz.Rect(block["bbox"])
        key = (int(block.get("xref") or 0), round(rect.x0, 2), round(rect.y0, 2), round(rect.x1, 2), round(rect.y1, 2))
        if key in seen:
            continue
        blob = _image_bytes(page.parent, block)
        if not blob:
            continue
        path = work_dir / f"image-{index:04d}.png"
        path.write_bytes(blob)
        path, rect = _clip_image(path, rect, page.rect)
        path = _rotate_image(path, _rotation(block.get("transform")))
        seen.add(key)
        results.append({"path": path, "rect": rect})

    try:
        infos = page.get_image_info(xrefs=True)
    except Exception:
        infos = []
    for index, info in enumerate(infos):
        if not info.get("bbox"):
            continue
        rect = fitz.Rect(info["bbox"])
        xref = int(info.get("xref") or 0)
        key = (xref, round(rect.x0, 2), round(rect.y0, 2), round(rect.x1, 2), round(rect.y1, 2))
        if key in seen or not xref:
            continue
        try:
            path = work_dir / f"xref-{index:04d}.png"
            path.write_bytes(page.parent.extract_image(xref)["image"])
            path, rect = _clip_image(path, rect, page.rect)
            path = _rotate_image(path, _rotation(info.get("transform")))
        except Exception:
            continue
        seen.add(key)
        results.append({"path": path, "rect": rect})

    results.sort(key=lambda x: (round(x["rect"].y0, 2), round(x["rect"].x0, 2)))
    return results


def _render_text(cell, block, cursor_y):
    bbox = fitz.Rect(block["bbox"])
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Pt(max(0, bbox.x0))
    p.paragraph_format.space_before = Pt(max(0, bbox.y0 - cursor_y))
    p.paragraph_format.space_after = Pt(0)
    first = True
    bottom = bbox.y1
    for line in block.get("lines", []):
        if not first:
            p.add_run().add_break()
        first = False
        for span in line.get("spans", []):
            text = span.get("text", "")
            if text:
                run = p.add_run(text)
                _set_run_font(run, span)
        bottom = max(bottom, fitz.Rect(line.get("bbox", bbox)).y1)
    return bottom


def _render_page(doc, page, page_index, work_dir):
    if page_index:
        doc.add_section(WD_SECTION.NEW_PAGE)
    section = doc.sections[-1]
    section.page_width = Inches(page.rect.width / 72)
    section.page_height = Inches(page.rect.height / 72)
    for name in ("top_margin", "bottom_margin", "left_margin", "right_margin", "header_distance", "footer_distance"):
        setattr(section, name, Inches(0))

    page_width = Inches(page.rect.width / 72)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    _set_table_width(table, page_width)
    _remove_table_borders(table)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _remove_cell_margins(cell)

    blocks = page.get_text("dict", sort=False).get("blocks", [])
    objects = [(float(fitz.Rect(b["bbox"]).y0), 0, b) for b in blocks if b.get("type") == 0 and b.get("lines")]
    images = _extract_images(page, work_dir)
    objects += [(float(i["rect"].y0), 1, i) for i in images]
    objects.sort(key=lambda x: (round(x[0], 2), x[1]))

    cursor_y = 0.0
    for _, kind, obj in objects:
        if kind == 0:
            cursor_y = _render_text(cell, obj, cursor_y)
        else:
            _add_anchor(cell, obj["path"], obj["rect"])
            cursor_y = max(cursor_y, obj["rect"].y1)


def convert_pdf_to_docx(source_pdf: Path, output: Path, progress=None):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    with fitz.open(source_pdf) as pdf, tempfile.TemporaryDirectory(prefix="pdftoword-v4-") as tmp:
        total = len(pdf)
        for index, page in enumerate(pdf):
            if progress:
                progress({"stage": "render", "percent": 10 + int(index / max(total, 1) * 80), "message": f"V4 正在解析第 {index + 1}/{total} 页…"})
            _render_page(doc, page, index, Path(tmp) / f"page-{index + 1}")
    if progress:
        progress({"stage": "save", "percent": 95, "message": "正在生成 DOCX…"})
    doc.save(output)
